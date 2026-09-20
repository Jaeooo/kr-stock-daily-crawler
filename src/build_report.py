"""티커리스트.xlsx(코드+종목명)를 읽어 네이버 증권 데이터로 채운 새 엑셀을 생성한다."""

from __future__ import annotations

import csv
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

import naver_client

SRC_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SRC_DIR.parent
DEFAULT_TICKER_LIST = PROJECT_ROOT / "티커리스트.xlsx"
DEFAULT_OUTPUT_DIR = PROJECT_ROOT / "output"

MAX_WORKERS = 6

COLUMN_WIDTHS = {
    "A": 4.55078125,
    "B": 19.8671875,
    "C": 7.84375,
    "D": 4.046875,
    "E": 5.56640625,
    "F": 10.125,
    "G": 5.56640625,
}

PLACEHOLDER = "-"


def get_target_date() -> date:
    """조회 대상 날짜. 지금은 항상 오늘 날짜를 쓰고, 나중에 GUI에서 날짜를 받으면
    이 함수만 교체하면 된다."""
    return datetime.now().date()


def get_available_dates(ticker_list_path: Path = DEFAULT_TICKER_LIST) -> list[date]:
    """네이버 API가 실제로 데이터를 갖고 있는 날짜 목록(최근 5거래일, 최신순)을 반환한다.

    integration API는 항상 최근 5거래일만 주기 때문에, 대표로 티커리스트의
    첫 종목 하나만 조회해서 그 날짜들을 기준으로 삼는다."""
    tickers = load_ticker_list(ticker_list_path)
    if not tickers:
        return []
    reference_code, _name = tickers[0]
    deal_trend = naver_client.fetch_deal_trend(reference_code)

    dates: list[date] = []
    for row in deal_trend:
        bizdate = row.get("bizdate")
        if bizdate:
            dates.append(datetime.strptime(bizdate, "%Y%m%d").date())
    return dates


def load_ticker_list(path: Path) -> list[tuple[str, str]]:
    """티커리스트.xlsx(코드, 종목명)를 읽어 (코드, 종목명) 리스트로 반환한다.
    코드가 비어있는 행은 건너뛴다."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    rows: list[tuple[str, str]] = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue  # 헤더
        code, name = row[0], row[1]
        if not code or not name:
            continue
        rows.append((str(code).strip(), str(name).strip()))
    return rows


def _strip_percent(value: str) -> str:
    return value[:-1] if value.endswith("%") else value


def _fetch_row(code: str, bizdate: str) -> tuple[str, str, str, str, str]:
    """(종가, 기관, 외국인, 외국인소진율, 거래량) 튜플을 반환. 실패 시 전부 '-'."""
    deal_trend = naver_client.fetch_deal_trend(code)
    row = naver_client.find_row_for_date(deal_trend, bizdate)
    if row is None:
        return (PLACEHOLDER,) * 5

    return (
        row.get("closePrice", PLACEHOLDER),
        row.get("organPureBuyQuant", PLACEHOLDER),
        row.get("foreignerPureBuyQuant", PLACEHOLDER),
        _strip_percent(row.get("foreignerHoldRatio", PLACEHOLDER)),
        row.get("accumulatedTradingVolume", PLACEHOLDER),
    )


def build_workbook(tickers: list[tuple[str, str]], target: date) -> tuple[openpyxl.Workbook, list[str]]:
    bizdate = target.strftime("%Y%m%d")
    date_label = f"{target.month}/{target.day}"

    wb = openpyxl.Workbook()
    ws: Worksheet = wb.active
    ws.title = f"{target.month}-{target.day}"

    headers = [None, "종목", f"{date_label} 종가", "기관", "외국인", "외국인소진율", "거래량", "종목코드"]
    ws.append(headers)

    for col, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[col].width = width
    ws.column_dimensions["H"].width = 10

    results: dict[int, tuple[str, str, str, str, str]] = {}
    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
        future_to_idx = {
            executor.submit(_fetch_row, code, bizdate): idx
            for idx, (code, _name) in enumerate(tickers)
        }
        for future in as_completed(future_to_idx):
            idx = future_to_idx[future]
            results[idx] = future.result()

    failed_names: list[str] = []
    for idx, (code, name) in enumerate(tickers):
        row_values = results[idx]
        if all(v == PLACEHOLDER for v in row_values):
            failed_names.append(name)
        ws.append([idx + 1, name, *row_values, code])

    for row in ws.iter_rows(min_row=2, min_col=8, max_col=8):
        row[0].number_format = "@"  # 텍스트 서식 (앞자리 0 보존)

    return wb, failed_names


def _write_log_csv(path: Path, names: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["종목명"])
        for name in names:
            writer.writerow([name])


def run(ticker_list_path: Path = DEFAULT_TICKER_LIST, output_dir: Path = DEFAULT_OUTPUT_DIR, target: date | None = None) -> Path:
    if target is None:
        target = get_target_date()

    print(f"[build_report] 대상 날짜: {target.isoformat()}")
    tickers = load_ticker_list(ticker_list_path)
    print(f"[build_report] 종목 {len(tickers)}개 로드 완료 ({ticker_list_path})")

    wb, failed_names = build_workbook(tickers, target)

    output_dir.mkdir(parents=True, exist_ok=True)
    out_path = output_dir / f"종가{target.month}-{target.day}.xlsx"
    wb.save(out_path)
    print(f"[build_report] 결과 저장: {out_path}")

    if failed_names:
        failed_path = output_dir / "failed_codes.csv"
        _write_log_csv(failed_path, failed_names)
        print(f"[build_report] 조회 실패 종목 {len(failed_names)}개 -> {failed_path}")

    return out_path
