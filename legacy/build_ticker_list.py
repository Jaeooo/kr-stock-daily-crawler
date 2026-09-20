"""종목리스트.txt(이름)를 참고해서 티커리스트.xlsx(코드,이름)를 만든다.

입력을 종목명 대신 종목코드로 바꾸는 걸 검토하기 위한 1회성 변환 스크립트.
매칭 안 되는 이름은 코드 칸을 비워두고 그대로 남겨서, 손으로 확인/수정할 수 있게 한다.
코드 열은 앞자리 0이 날아가지 않도록 텍스트 서식으로 고정한다.

매칭 안 된 이름은 티커매칭실패.xlsx에 (이름이 가장 비슷한) 네이버 후보명과 함께 따로 뽑는다.
"""

from __future__ import annotations

import difflib
import sys
from pathlib import Path

import openpyxl

LEGACY_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = LEGACY_DIR.parent
sys.path.insert(0, str(PROJECT_ROOT / "src"))

import stock_master  # noqa: E402 (src를 sys.path에 넣은 뒤에 임포트)

STOCK_LIST_PATH = LEGACY_DIR / "종목리스트.txt"
OUTPUT_PATH = PROJECT_ROOT / "티커리스트.xlsx"
UNMATCHED_PATH = PROJECT_ROOT / "티커매칭실패.xlsx"


def load_stock_list(path: Path) -> list[str]:
    names: list[str] = []
    with path.open("r", encoding="utf-8-sig") as f:
        for line in f:
            name = line.strip()
            if name:
                names.append(name)
    return names


def main() -> None:
    names = load_stock_list(STOCK_LIST_PATH)
    name_to_code = stock_master.get_name_to_code_map()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "티커리스트"
    ws.append(["코드", "종목명"])
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 20

    unmatched_names: list[str] = []
    for name in names:
        code = name_to_code.get(name, "")
        if not code:
            unmatched_names.append(name)
        ws.append([code, name])

    for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
        row[0].number_format = "@"  # 텍스트 서식 (앞자리 0 보존)

    wb.save(OUTPUT_PATH)

    candidate_names = list(name_to_code.keys())
    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    ws2.title = "티커매칭실패"
    ws2.append(["종목명", "네이버 후보명(참고용)"])
    ws2.column_dimensions["A"].width = 20
    ws2.column_dimensions["B"].width = 20
    for name in unmatched_names:
        matches = difflib.get_close_matches(name, candidate_names, n=1, cutoff=0.6)
        ws2.append([name, matches[0] if matches else ""])
    wb2.save(UNMATCHED_PATH)

    print(f"[build_ticker_list] {len(names)}개 중 {len(names) - len(unmatched_names)}개 매칭, {len(unmatched_names)}개 미매칭")
    print(f"[build_ticker_list] 저장: {OUTPUT_PATH}")
    print(f"[build_ticker_list] 미매칭 목록: {UNMATCHED_PATH}")


if __name__ == "__main__":
    main()
